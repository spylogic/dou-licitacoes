# -*- coding: utf-8 -*-
"""
gerar_excel.py
Monta o arquivo .xlsx de saída a partir da lista de MateriaDOU retornada
pelo dou_scraper.

O arquivo tem duas abas:
  - "Panorama": um resumo geral (total de matérias, contagem por tipo de
    ato, contagem por órgão principal) — útil para uma visão rápida do dia.
  - "Relação": a lista detalhada, com as colunas DATA / TÍTULO / OBJETO /
    RESUMO BREVE (mais ÓRGÃO e TIPO DE ATO, que ajudam bastante e não
    quebram o pedido original), com cabeçalho formatado, largura de coluna
    ajustada e AutoFilter.
"""
from __future__ import annotations

from collections import Counter
from datetime import date as date_cls
from io import BytesIO
from typing import List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from dou_scraper import MateriaDOU

# Paleta navy/dourado, consistente com o tema visual do painel (index.html)
COR_NAVY = "0A2547"
COR_NAVY_CLARO = "13345E"
COR_DOURADO = "C9A227"
COR_TEXTO_CLARO = "FFFFFF"
COR_LINHA_ALT = "F4F1E8"


def _cabecalho_estilo():
    return {
        "font": Font(bold=True, color=COR_TEXTO_CLARO, size=11),
        "fill": PatternFill(start_color=COR_NAVY, end_color=COR_NAVY, fill_type="solid"),
        "alignment": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _aplicar_bordas_finas(ws: Worksheet, linha_inicio: int, linha_fim: int, col_inicio: int, col_fim: int):
    borda = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )
    for row in ws.iter_rows(min_row=linha_inicio, max_row=linha_fim, min_col=col_inicio, max_col=col_fim):
        for cell in row:
            cell.border = borda


def _montar_panorama(ws: Worksheet, materias: List[MateriaDOU], data_consulta: date_cls, so_licitacao: bool):
    estilo_cab = _cabecalho_estilo()

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14

    ws.merge_cells("A1:B1")
    titulo_cell = ws["A1"]
    titulo_cell.value = "Achador de Licitações — Panorama do DOU"
    titulo_cell.font = Font(bold=True, color=COR_TEXTO_CLARO, size=14)
    titulo_cell.fill = PatternFill(start_color=COR_NAVY, end_color=COR_NAVY, fill_type="solid")
    titulo_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    linha = 3
    ws.cell(row=linha, column=1, value="Data consultada").font = Font(bold=True)
    ws.cell(row=linha, column=2, value=data_consulta.strftime("%d/%m/%Y"))
    linha += 1

    ws.cell(row=linha, column=1, value="Seção").font = Font(bold=True)
    ws.cell(row=linha, column=2, value="Seção 3 — Contratos, Editais e Avisos")
    linha += 1

    ws.cell(row=linha, column=1, value="Filtro aplicado").font = Font(bold=True)
    ws.cell(row=linha, column=2, value="Somente licitação" if so_licitacao else "Todas as matérias da Seção 3")
    linha += 1

    ws.cell(row=linha, column=1, value="Total de matérias listadas").font = Font(bold=True)
    ws.cell(row=linha, column=2, value=len(materias)).font = Font(bold=True, color=COR_NAVY)
    linha += 2

    # Contagem por tipo de ato
    linha_inicio_tipos = linha
    cab = ws.cell(row=linha, column=1, value="Tipo de ato")
    cab2 = ws.cell(row=linha, column=2, value="Quantidade")
    for c in (cab, cab2):
        c.font = estilo_cab["font"]
        c.fill = estilo_cab["fill"]
        c.alignment = estilo_cab["alignment"]
    linha += 1

    contagem_tipo = Counter(m.tipo_ato or "(não informado)" for m in materias)
    for tipo, qtd in contagem_tipo.most_common():
        ws.cell(row=linha, column=1, value=tipo)
        ws.cell(row=linha, column=2, value=qtd)
        if (linha - linha_inicio_tipos) % 2 == 0:
            for col in (1, 2):
                ws.cell(row=linha, column=col).fill = PatternFill(
                    start_color=COR_LINHA_ALT, end_color=COR_LINHA_ALT, fill_type="solid"
                )
        linha += 1

    if contagem_tipo:
        _aplicar_bordas_finas(ws, linha_inicio_tipos, linha - 1, 1, 2)

    linha += 1

    # Contagem por órgão principal (primeiro nível da hierarquia)
    linha_inicio_orgaos = linha
    cab = ws.cell(row=linha, column=1, value="Órgão principal")
    cab2 = ws.cell(row=linha, column=2, value="Quantidade")
    for c in (cab, cab2):
        c.font = estilo_cab["font"]
        c.fill = estilo_cab["fill"]
        c.alignment = estilo_cab["alignment"]
    linha += 1

    contagem_orgao = Counter(
        (m.orgao.split("/")[0] if m.orgao else "(não informado)") for m in materias
    )
    for orgao, qtd in contagem_orgao.most_common(20):
        ws.cell(row=linha, column=1, value=orgao)
        ws.cell(row=linha, column=2, value=qtd)
        if (linha - linha_inicio_orgaos) % 2 == 0:
            for col in (1, 2):
                ws.cell(row=linha, column=col).fill = PatternFill(
                    start_color=COR_LINHA_ALT, end_color=COR_LINHA_ALT, fill_type="solid"
                )
        linha += 1

    if contagem_orgao:
        _aplicar_bordas_finas(ws, linha_inicio_orgaos, linha - 1, 1, 2)

    ws.freeze_panes = "A2"


def _montar_relacao(ws: Worksheet, materias: List[MateriaDOU]):
    colunas = [
        ("DATA", 12),
        ("TÍTULO", 38),
        ("OBJETO", 55),
        ("RESUMO BREVE", 55),
        ("ÓRGÃO", 45),
        ("TIPO DE ATO", 28),
        ("LINK", 40),
    ]

    estilo_cab = _cabecalho_estilo()
    for idx, (nome, largura) in enumerate(colunas, start=1):
        cell = ws.cell(row=1, column=idx, value=nome)
        cell.font = estilo_cab["font"]
        cell.fill = estilo_cab["fill"]
        cell.alignment = estilo_cab["alignment"]
        ws.column_dimensions[get_column_letter(idx)].width = largura

    ws.row_dimensions[1].height = 22

    for i, m in enumerate(materias, start=2):
        ws.cell(row=i, column=1, value=m.data)
        ws.cell(row=i, column=2, value=m.titulo)
        ws.cell(row=i, column=3, value=m.objeto)
        ws.cell(row=i, column=4, value=m.resumo_breve)
        ws.cell(row=i, column=5, value=m.orgao)
        ws.cell(row=i, column=6, value=m.tipo_ato)
        if m.url:
            cell_link = ws.cell(row=i, column=7, value=m.url)
            cell_link.hyperlink = m.url
            cell_link.font = Font(color="1155CC", underline="single")

        for col in range(1, len(colunas) + 1):
            ws.cell(row=i, column=col).alignment = Alignment(
                vertical="top", wrap_text=(col in (2, 3, 4, 5))
            )

        if i % 2 == 0:
            for col in range(1, len(colunas) + 1):
                cell = ws.cell(row=i, column=col)
                if cell.fill.start_color.rgb in (None, "00000000"):
                    cell.fill = PatternFill(
                        start_color=COR_LINHA_ALT, end_color=COR_LINHA_ALT, fill_type="solid"
                    )

    ultima_linha = max(len(materias) + 1, 1)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{ultima_linha}"
    ws.freeze_panes = "A2"

    if materias:
        _aplicar_bordas_finas(ws, 1, ultima_linha, 1, len(colunas))


def gerar_excel_bytes(
    materias: List[MateriaDOU],
    data_consulta: date_cls,
    so_licitacao: bool = False,
) -> bytes:
    """Gera o arquivo .xlsx em memória e retorna os bytes prontos para download."""
    wb = Workbook()

    ws_panorama = wb.active
    ws_panorama.title = "Panorama"
    _montar_panorama(ws_panorama, materias, data_consulta, so_licitacao)

    ws_relacao = wb.create_sheet("Relação")
    _montar_relacao(ws_relacao, materias)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    # Teste manual rápido usando o fixture de testes (não depende de internet)
    import json
    import os

    fixture_path = os.path.join(os.path.dirname(__file__), "tests", "fixture_dou_params.json")
    with open(fixture_path, encoding="utf-8") as f:
        fixture = json.load(f)

    from dou_scraper import _item_para_materia

    materias = [_item_para_materia(item) for item in fixture["jsonArray"]]
    conteudo = gerar_excel_bytes(materias, date_cls(2026, 9, 9))
    saida = os.path.join(os.path.dirname(__file__), "tests", "saida_teste.xlsx")
    with open(saida, "wb") as f:
        f.write(conteudo)
    print(f"Arquivo de teste gerado em: {saida} ({len(conteudo)} bytes)")
